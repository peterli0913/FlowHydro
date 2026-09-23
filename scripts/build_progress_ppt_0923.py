"""
Build a 4-slide biweekly progress deck (16:9) for 9–23 Sep 2026,
for the leadership review.  Same chrome as the 0826 / 0909 decks:
English primary, Chinese secondary.

Sources
-------
- Meeting Minutes Tracking Record 0922.xlsx (sheet 0922 vs 0907)
- 22 Sep technical session with Keith
  (文字记录：20260922_161033)

What this period actually moved
-------------------------------
No tracker line moved to YES (still 62 closed / 44 pending / 29 open).
CFCT issued the 3D on 21 Sep.  The 22 Sep meeting turned that model
into layout rules and moved the main-equipment gate onto nozzle
orientation and the HTF runs.

Omitted
-------
Travel-approval discussion, why a window might move, and any cost
figure that was only guessed in the room.  The early-November visit
window is stated as still being checked, not as a decision.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

FONT_EN = "Segoe UI"
FONT_ZH = "Microsoft YaHei"

BODY = "1E2B3A"
SUB = "8A97AB"
FOOT = "9AA6B8"
LINE = "E3E8F0"

PERIOD = "Bi-weekly Progress 2026-09-09 ~ 09-23"
TOTAL = "04"
OUT = "Progress_Report_0923.pptx"


def rgb(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _run(p, text, *, font=FONT_EN, size=18, bold=False, color=BODY, italic=False):
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = rgb(color)
    rPr = r._r.get_or_add_rPr()
    rPr.append(rPr.makeelement(qn("a:ea"), {"typeface": FONT_ZH}))
    return r


def box(slide, kind, x, y, w, h, *, fill=None, line=None, lw=1.0):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def add_footer(slide, no):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(7.05),
                                  Inches(7.2), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    _run(p, "Asymchem · Sandwich Site   |   Continuous Hydrogenation Skid",
         size=9, color=FOOT)
    tb = slide.shapes.add_textbox(Inches(7.6), Inches(7.05),
                                  Inches(5.1), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    _run(p, f"{PERIOD}    ·    {no} / {TOTAL}", size=9, color=FOOT)


def add_header(slide, no, title, sub, color):
    box(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.16, fill=color)
    box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.55, 0.38, 0.92, 0.92, fill=color)
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(0.38),
                                  Inches(0.92), Inches(0.92))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, no, size=32, bold=True, color="FFFFFF")

    tb = slide.shapes.add_textbox(Inches(1.68), Inches(0.40),
                                  Inches(11.0), Inches(0.88))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    _run(p, title, size=28, bold=True, color=color)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(1)
    _run(p2, sub, font=FONT_ZH, size=14, color=SUB)

    box(slide, MSO_SHAPE.RECTANGLE, 0.55, 1.46, 12.23, 0.015, fill=LINE)


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_bullet_block(slide, y, color, en, cn, *, small=False, subs=None, gap=0.14):
    subs = subs or []
    en_sz = 14 if small else 16
    cn_sz = 10.5 if small else 11.5
    dot_sz = 14 if small else 16
    h_box = 0.48 + 0.30 * len(subs)
    tb = slide.shapes.add_textbox(Inches(0.62), Inches(y),
                                  Inches(12.1), Inches(h_box))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    p = tf.paragraphs[0]
    _run(p, "●  ", size=dot_sz, bold=True, color=(SUB if small else color))
    _run(p, en, size=en_sz, bold=(not small),
         color=(BODY if not small else "5A6675"))
    p2 = tf.add_paragraph()
    p2.space_before = Pt(1)
    _run(p2, "     " + cn, font=FONT_ZH, size=cn_sz, color=SUB)
    for se, sc in subs:
        ps = tf.add_paragraph()
        ps.space_before = Pt(2)
        _run(ps, "        –  ", size=12, color="B7C0CF")
        _run(ps, se, size=13, color="3C4A5A")
        pc = tf.add_paragraph()
        _run(pc, "             " + sc, font=FONT_ZH, size=10, color=SUB)
    return y + 0.46 + 0.30 * len(subs) + gap


def slide_01(prs):
    color = "2563EB"
    s = blank_slide(prs)
    add_header(s, "01", "Overall Progress", "整体推进", color)
    y = 1.64
    y = add_bullet_block(
        s, y, color,
        "21 Sep — first 3D on the 8 Sep direction, issued by email.  "
        "22 Sep — reviewed live with Keith.",
        "9 月 21 日发出按 9 月 8 日方向排的第一版 3D；9 月 22 日与 Keith 对照模型过完。",
        gap=0.12,
    )
    y = add_bullet_block(
        s, y, color,
        "The maintenance principle is unchanged: pull the front module out.  "
        "No central walkway.  Shipping height stays 2.5 m.",
        "维护原则未变：前撬拉出检修，不留中间通道，运输高度仍为 2.5 m。",
        gap=0.12,
    )
    y = add_bullet_block(
        s, y, color,
        "Tracker: 62 closed, 44 pending, 29 open.  "
        "Nothing moved to closed this fortnight.",
        "跟踪表：已关闭 62，待定 44，未关闭 29。本双周没有新的关闭项。",
        subs=[
            ("The movement is on the model and on three rules Keith set "
             "on 22 Sep — not on line-by-line closure.",
             "本期推进在模型和 22 日 Keith 定下的三条规则上，不在逐条关项上。"),
        ],
        gap=0.12,
    )
    y = add_bullet_block(
        s, y, color,
        "Critical path has narrowed: nozzle direction and the HTF runs, "
        "then Keith can release the 18 Aug equipment drawings.",
        "关键路径收窄为：先定管嘴朝向和 HTF 走向，Keith 才能放行 8 月 18 日的主设备图。",
        gap=0.08,
    )
    add_footer(s, "01")
    return y


def slide_02(prs):
    color = "16A34A"
    s = blank_slide(prs)
    add_header(s, "02", "What 22 Sep Settled", "9 月 22 日已说清的技术规则", color)
    y = 1.62
    y = add_bullet_block(
        s, y, color,
        "TCU — the wall penetration is fixed.  Ten pipes may stack after "
        "the wall.  The five TCU units follow our nozzles.",
        "TCU：穿墙位置已定。10 根管过墙后可以叠成两层。5 台 TCU 跟着我方管口走。",
        subs=[
            ("One vent on the SW side of the wall, not ten vents on the skid.  "
             "Wall plate is flanged both sides — no site weld in the gap.",
             "排气只留一个点，放在墙的 SW 侧，不在撬上放 10 个排气阀。"
             "穿墙板两侧法兰对接，窄处不现场焊接。"),
        ],
        gap=0.08,
    )
    y = add_bullet_block(
        s, y, color,
        "Layout — two options will be drawn and compared.  Not a freeze yet.",
        "布置先出两版对比，这次会上没有冻结。",
        subs=[
            ("Fan’s option: swap SE01 / CU01 with the rear valves, so the "
             "valves stay in front when the module is pulled.",
             "范双双的方案：SE01 / CU01 与后部阀门对调，前撬拉出后阀门仍在前面。"),
            ("Keith’s option: condenser slightly forward; separator turned "
             "to the wall so the level gauge withdraws from the end.  "
             "The hole in the middle takes fully welded HTF.",
             "Keith 的方案：冷凝器略向前；分离器靠墙转约 90°，液位计从端部抽出。"
             "中间空档走全焊接 HTF。"),
        ],
        gap=0.08,
    )
    y = add_bullet_block(
        s, y, color,
        "Reactor 3 HTF nozzle turns toward the wall riser.  "
        "The pipe no longer runs out to the skid end and doubles back.",
        "反应器 3 的 HTF 管嘴改朝穿墙立管。管子不再先跑到撬端再折回。",
        gap=0.08,
    )
    y = add_bullet_block(
        s, y, color,
        "Drains and CIP — valve at the tee, hose at the front.  "
        "PSV-025 and PSV-019 stay separate and move to the skid edge.",
        "排净和 CIP：阀紧贴三通，软管接口在撬前。PSV-025 与 PSV-019 不合并，挪到撬边。",
        subs=[
            ("Unreachable valves take a compact Habonim actuated ball valve, "
             "operated from the end.  A 300–400 mm reach can stay manual.",
             "够不到的阀用 Habonim 小型气动球阀，在撬端操作。"
             "伸手 300–400 mm 能到的可以保持手动。"),
        ],
        gap=0.06,
    )
    add_footer(s, "02")
    return y


def slide_03(prs):
    color = "DC2626"
    s = blank_slide(prs)
    add_header(s, "03", "What Changes the Plan", "相对 9 月 9 日计划的变化", color)
    y = 1.62
    y = add_bullet_block(
        s, y, color,
        "Main-equipment enquiry now waits on nozzle orientation and the "
        "HTF runs — not on a full drawing sign-off, and not on the UK visit.",
        "主设备询价改为等管嘴朝向和 HTF 走向。不等整套图纸签完，也不等赴英。",
        subs=[
            ("9 Sep allowed the enquiry once the reactor specification was "
             "confirmed.  On 22 Sep Keith would not comment on the 18 Aug "
             "drawings until those nozzles are shown to fit.",
             "9 月 9 日的口径是反应柱规格确认后即可询价。"
             "22 日 Keith 明确：管嘴方向和 HTF 没在模型里站住之前，他不审 8 月 18 日的图。"),
        ],
        gap=0.10,
    )
    y = add_bullet_block(
        s, y, color,
        "Charging stairs — SW supplies them in the UK, in stainless steel, "
        "separate from the main-equipment order.",
        "加料梯由 SW 在英国供货，不锈钢，不并进主设备采购。",
        subs=[
            ("A catalogue aluminium stair is not acceptable.  "
             "China will not design or fabricate this item.  "
             "Target height is under 1.4 m.",
             "不用成品铝梯。中方不设计、不制造。高度目标低于 1.4 m。"),
        ],
        gap=0.10,
    )
    y = add_bullet_block(
        s, y, color,
        "12–22 October is still the proposed visit.  It is not confirmed.  "
        "An early-November window is being checked.",
        "10 月 12–22 日仍是建议行程，尚未确认。11 月初的备选窗口正在核实。",
        subs=[
            ("The model work this week does not wait for that confirmation.",
             "本周的模型工作不等这个日期。"),
        ],
        gap=0.10,
    )
    y = add_bullet_block(
        s, y, color,
        "Still open, and not a layout choice for this meeting: "
        "width cannot drop 100–200 mm (pump and BT01).  "
        "Cleaning-mode high-level bypass is in the manual, not yet agreed.",
        "本次不需要领导选型。宽度受泵和 BT01 限制，减 100–200 mm 没有意义。"
        "清洗模式旁路高液位联锁已写入手册，Keith 尚未认可。",
        small=True, gap=0.06,
    )
    add_footer(s, "03")
    return y


def slide_04(prs):
    color = "0E9AA7"
    s = blank_slide(prs)
    add_header(s, "04", "Next Steps", "下一步", color)
    y = 1.62
    steps = [
        ("Fan — draw both layouts, turn the reactor-3 nozzle toward the "
         "wall riser, and place drain / CIP valves at the tee.",
         "范双双：两版布置都画出来；反应器 3 管嘴朝向穿墙立管；排净 / CIP 阀放到三通处。"),
        ("Gao Yu — Habonim ball valve and actuator envelope to Fan, "
         "and whether it can be bought in China.",
         "高宇：把 Habonim 球阀和气动头的外形给范双双，并确认国内能否买到。"),
        ("Meng — which signals stay on the pull-out module, and which "
         "still leave it.  Cable count waits on the junction-box size.",
         "孟德智：可拉出模块上哪些信号留在本撬、哪些还要外引。电缆数量等接线箱尺寸。"),
        ("Keith — send the NWD model; junction-box feedback by "
         "25 or 28 Sep; valve tidy-up once the layout is close.",
         "Keith：发 NWD；接线箱 9 月 25 日或 28 日反馈；布置接近后再精简阀门。"),
        ("Zhao — hold the 18 Aug equipment issue until the nozzles and "
         "HTF runs are accepted.  CE document set continues in parallel.",
         "赵子亮：管嘴和 HTF 被接受前，不按 8 月 18 日图提采购。CE 资料并行准备。"),
        ("Confirm the visit window with Keith.  Next technical session: "
         "23 Sep, on the same model.",
         "与 Keith 确认赴英窗口。下一次技术会是 9 月 23 日，仍对着这版模型。"),
    ]
    for en, cn in steps:
        y = add_bullet_block(s, y, color, en, cn, gap=0.08)
    add_footer(s, "04")
    return y


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    ys = [slide_01(prs), slide_02(prs), slide_03(prs), slide_04(prs)]
    prs.save(OUT)
    print("Saved:", OUT, "end-y:", [round(y, 2) for y in ys])


if __name__ == "__main__":
    main()
