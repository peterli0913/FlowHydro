"""
English working-meeting deck for the 8 Sep 2026 discussion with Keith.

Topics are the tracker items locked in the 7 Sep internal prep:
gasket vendor result, 3D / three-cavity / height / control panel,
separator back-pressure valve vs 270 °C, PSV if DN50, and tracker close-out.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

FONT = "Calibri"
NAVY = "00366A"
ORANGE = "E36C09"
TEAL = "0E9AA7"
RED = "B42318"
GREEN = "1B7A4E"
AMBER = "B45309"
BODY = "1E2B3A"
MUTED = "5B6B7C"
LINE = "D7DEE8"
FOOT = "8A97AB"
WHITE = "FFFFFF"
ICE = "F4F7FB"
AMBER_BG = "FFF6E8"
RED_BG = "FDECEC"
GREEN_BG = "EAF6EF"
TEAL_BG = "E7F6F7"
NAVY_BG = "E8EEF6"

W, H = 13.333, 7.5
OUT = "Keith_Discussion_0908.pptx"


def rgb(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def run(p, text, *, size=16, bold=False, color=BODY, font=FONT, italic=False):
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = rgb(color)
    rPr = r._r.get_or_add_rPr()
    rPr.append(rPr.makeelement(qn("a:ea"), {"typeface": font}))
    rPr.append(rPr.makeelement(qn("a:cs"), {"typeface": font}))
    return r


def shape(slide, kind, x, y, w, h, *, fill=None, line=None, lw=1.0):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def tb(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    return tf


def p0(tf, text, **kw):
    p = tf.paragraphs[0]
    run(p, text, **kw)
    return p


def addp(tf, text, *, before=0, **kw):
    p = tf.add_paragraph()
    if before:
        p.space_before = Pt(before)
    run(p, text, **kw)
    return p


def footer(slide, page, total=7):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 7.28, W, 0.22, fill=NAVY)
    tf = tb(slide, 0.45, 7.28, 8.2, 0.22, anchor=MSO_ANCHOR.MIDDLE)
    p0(tf, "Asymchem  ·  Sandwich continuous hydrogenation  ·  Discussion with Keith",
       size=10, color=WHITE)
    tf = tb(slide, 10.4, 7.28, 2.5, 0.22, anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run(p, f"8 Sep 2026    {page} / {total}", size=10, color=WHITE)


def header(slide, kicker, title, accent=NAVY):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, W, 0.10, fill=accent)
    tf = tb(slide, 0.45, 0.22, 12.4, 0.28)
    p0(tf, kicker, size=12, bold=True, color=accent)
    tf = tb(slide, 0.45, 0.48, 12.4, 0.48)
    p0(tf, title, size=26, bold=True, color=NAVY)


def ask_bar(slide, text, *, y=6.58, fill=NAVY, label="ASK"):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, y, 12.4, 0.58,
          fill=fill)
    tf = tb(slide, 0.65, y, 12.0, 0.58, anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    run(p, f"{label}   ", size=13, bold=True, color="FDE68A")
    run(p, text, size=14, bold=True, color=WHITE)


def slide_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=NAVY)
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 0.18, H, fill=ORANGE)
    tf = tb(s, 0.70, 1.55, 12.0, 0.35)
    p0(tf, "SANDWICH  ·  CONTINUOUS HYDROGENATION", size=14, bold=True, color="9DB4D0")
    tf = tb(s, 0.70, 2.00, 12.0, 1.50)
    p0(tf, "Technical discussion with Keith", size=36, bold=True, color=WHITE)
    addp(tf, "Tracker close-out and one new instrument risk", size=20, color="D5E2F0", before=8)
    tf = tb(s, 0.70, 4.00, 12.0, 0.90)
    p0(tf, "8 September 2026", size=16, bold=True, color="F6C089")
    addp(tf, "Internal working session  ·  not a new HAZOP  ·  close what we can online",
         size=14, color="B8C7DA", before=4)

    items = [
        ("01", "Gasket"),
        ("02", "3D layout"),
        ("03", "PIC-028 vs 270 °C"),
        ("04", "PSV if DN50"),
        ("05", "Tracker close-out"),
    ]
    x = 0.70
    for no, name in items:
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 5.55, 2.30, 0.95,
              fill="0A2A56")
        tf = tb(s, x + 0.12, 5.55, 2.06, 0.95, anchor=MSO_ANCHOR.MIDDLE)
        p0(tf, no, size=12, bold=True, color="F6C089")
        addp(tf, name, size=13, bold=True, color=WHITE, before=2)
        x += 2.45
    return s


def slide_agenda(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=WHITE)
    header(s, "AGENDA", "What we need to cover today")

    rows = [
        ("01", GREEN, "INFORM", "High-pressure gasket",
         "Vendor can supply a 316L metal-ring gasket if hardness limits are written into the spec."),
        ("02", TEAL, "DISCUSS", "3D layout — three-cavity, height, control panel",
         "A draft based on Keith’s three-cavity idea is ready. Close what we can before the site week."),
        ("03", RED, "DECIDE", "Separator vapour-outlet valve vs 270 °C",
         "Bronkhorst electronics may fail at 270 °C. No qualified alternative yet. This is the new risk."),
        ("04", AMBER, "CLARIFY", "Relief-valve size if the reactor is DN50",
         "DN80 and DN65 already share the same orifice. Confirm whether DN50 still can."),
        ("05", NAVY, "CONFIRM", "Tracker lines already answered",
         "Cleaning, operating manual and control description — please confirm which lines can close."),
    ]
    y = 1.20
    for no, color, tag, title, blurb in rows:
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, y, 12.4, 0.98,
              fill=ICE, line=LINE, lw=1.0)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.60, y + 0.24, 0.70, 0.50, fill=color)
        tf = tb(s, 0.60, y + 0.24, 0.70, 0.50, anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run(p, no, size=16, bold=True, color=WHITE)

        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.50, y + 0.18, 1.35, 0.28, fill=color)
        tf = tb(s, 1.50, y + 0.18, 1.35, 0.28, anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run(p, tag, size=10, bold=True, color=WHITE)

        tf = tb(s, 3.00, y + 0.12, 9.55, 0.74)
        p0(tf, title, size=16, bold=True, color=NAVY)
        addp(tf, blurb, size=13, color=MUTED, before=3)
        y += 1.07

    footer(s, "02")
    return s


def slide_gasket(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=WHITE)
    header(s, "01  ·  INFORM / CONFIRM", "High-pressure gasket — vendor result")

    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 1.18, 6.05, 5.00,
          fill=ICE, line=LINE)
    tf = tb(s, 0.70, 1.35, 5.55, 0.40)
    p0(tf, "Already agreed on 25 Aug", size=14, bold=True, color=TEAL)
    points = [
        "Primary route: 316L octagonal metal-ring gasket.",
        "Tantalum remains the expensive fallback.",
        "PTFE spiral-wound is out — code does not allow PTFE at 270 °C.",
        "Keith accepts gasket damage after a 270 °C relief event.",
        "The gasket must still seal during relief. A leak into the module is an ignition risk.",
    ]
    tf = tb(s, 0.70, 1.85, 5.55, 4.20)
    for i, t in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(8 if i else 0)
        run(p, "●  ", size=14, color=TEAL)
        run(p, t, size=14, color=BODY)

    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.70, 1.18, 6.15, 5.00,
          fill=GREEN_BG, line="B7E0C5")
    tf = tb(s, 6.95, 1.35, 5.70, 0.40)
    p0(tf, "Vendor position now", size=14, bold=True, color=GREEN)
    specs = [
        ("Gasket hardness", "< 150 HB"),
        ("Pipe-flange hardness", "> 160 HB"),
        ("Evidence", "Hardness test report for gasket and flange"),
        ("Check", "Equipment vendor re-inspects before assembly"),
        ("Scope", "Vessel joints and piping gaskets from one supplier"),
    ]
    y = 1.85
    for label, val in specs:
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.95, y, 5.70, 0.72, fill=WHITE)
        tf = tb(s, 7.10, y, 5.40, 0.72, anchor=MSO_ANCHOR.MIDDLE)
        p0(tf, label, size=11, color=MUTED)
        addp(tf, val, size=15, bold=True, color=NAVY)
        y += 0.80

    ask_bar(s, "Please confirm SW accepts this hardness specification and single-vendor scope.")
    footer(s, "03")
    return s


def slide_layout(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=WHITE)
    header(s, "02  ·  DISCUSS", "3D layout — close what we can online")

    cards = [
        (TEAL, TEAL_BG, "Three-cavity draft",
         "We have drawn a layout from Keith’s three-cavity / split-module idea.",
         "Walk the draft today. Agree what is acceptable so the site week is not used to rediscover the same points."),
        (AMBER, AMBER_BG, "Overall height",
         "The skid height has already been reduced on our side.",
         "Confirm the remaining height envelope with the split-module and goods-lift constraints."),
        (RED, RED_BG, "Control panel",
         "Still the hard item: shower clash, ATEX split, cabling, and whether the panel can hinge or ship loose.",
         "Need Keith’s instrument-to-module list before the cabinet can be frozen."),
    ]
    x = 0.45
    for color, bg, title, now, need in cards:
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.20, 4.00, 5.00,
              fill=bg, line=LINE)
        shape(s, MSO_SHAPE.RECTANGLE, x, 1.20, 4.00, 0.10, fill=color)
        tf = tb(s, x + 0.22, 1.45, 3.56, 0.70)
        p0(tf, title, size=18, bold=True, color=NAVY)

        tf = tb(s, x + 0.22, 2.20, 3.56, 1.70)
        p0(tf, "WHERE WE ARE", size=11, bold=True, color=color)
        addp(tf, now, size=14, color=BODY, before=6)

        tf = tb(s, x + 0.22, 4.05, 3.56, 2.00)
        p0(tf, "TODAY", size=11, bold=True, color=color)
        addp(tf, need, size=14, color=BODY, before=6)
        x += 4.15

    ask_bar(s, "Review the three-cavity draft now. Leave only items that truly need the Sandwich room.")
    footer(s, "04")
    return s


def slide_valve(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=WHITE)
    header(s, "03  ·  DECIDE", "PIC-028  ·  vapour-outlet back-pressure vs 270 °C",
           accent=RED)

    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 1.18, 12.4, 0.72,
          fill=RED_BG, line="F0B4B0")
    tf = tb(s, 0.65, 1.18, 12.0, 0.72, anchor=MSO_ANCHOR.MIDDLE)
    p0(tf, "New risk.  Gas–liquid separator vapour outlet.  Currently a Bronkhorst electronic back-pressure valve.",
       size=15, bold=True, color=RED)

    blocks = [
        (NAVY, ICE, "Duty",
         ["Tag: PIC-028 / PCV-028 on the separator vapour line.",
          "Wide pressure range, low flow — this is why the electronic valve was chosen.",
          "System design temperature remains 270 °C after the 25 Aug discussion."]),
        (AMBER, AMBER_BG, "What we found",
         ["Bronkhorst: electronics may fail instantly at 270 °C.",
          "Pneumatic substitutes checked so far cannot cover the ΔP and the low flow together.",
          "Only one vendor asked. Procurement is expanding the list before a formal enquiry."]),
        (RED, RED_BG, "Why this is not “replace the instrument”",
         ["If the valve dies in the relief case it may lose back-pressure.",
          "Inventory then dumps downstream / to vent — not a leak into the module, but not a safe silent failure either.",
          "We will not leave this as an unstated hazard."]),
    ]
    x = 0.45
    for color, bg, title, lines in blocks:
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.08, 4.00, 3.28,
              fill=bg, line=LINE)
        tf = tb(s, x + 0.18, 2.22, 3.64, 0.36)
        p0(tf, title, size=14, bold=True, color=color)
        tf = tb(s, x + 0.18, 2.62, 3.64, 2.58)
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(8 if i else 0)
            run(p, "●  ", size=13, color=color)
            run(p, line, size=13, color=BODY)
        x += 4.15

    ask_bar(
        s,
        "Must this tag survive 270 °C?  If no qualified valve exists, what architecture will SW accept?",
        fill=RED,
        label="DECIDE",
    )
    footer(s, "05")
    return s


def slide_psv(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=WHITE)
    header(s, "04  ·  CLARIFY", "Relief-valve size if the reactor moves to DN50")

    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 1.20, 8.15, 5.15,
          fill=ICE, line=LINE)
    tf = tb(s, 0.70, 1.40, 7.70, 0.40)
    p0(tf, "Working assumption", size=14, bold=True, color=NAVY)
    points = [
        "Reactor path: DN80 is dropped. Working case is DN65. DN50 is the backup if the September trial requires it.",
        "R&D trial on a DN65 column is now expected from about 18 Sep, with data by end of September.",
        "The current relief-valve selection is already the same orifice for DN80 and DN65.",
        "A DN50 case would cut heat generation further. We would rather not re-open the PSV file unless SW requires it.",
        "PFD and mass balance stay unchanged — only equipment datasheets move if the bore changes.",
    ]
    tf = tb(s, 0.70, 1.95, 7.70, 4.10)
    for i, t in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(10 if i else 0)
        run(p, "●  ", size=15, color=AMBER)
        run(p, t, size=15, color=BODY)

    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.80, 1.20, 4.05, 5.15,
          fill=AMBER_BG, line="F2D19A")
    tf = tb(s, 9.05, 1.45, 3.55, 0.40)
    p0(tf, "Please confirm", size=14, bold=True, color=AMBER)
    asks = [
        "If the reactor is DN65 — keep the current PSV.",
        "If the reactor is DN50 — may we keep that same orifice?",
        "Or must the relief case be re-run before HAZOP?",
    ]
    tf = tb(s, 9.05, 2.05, 3.55, 3.90)
    for i, t in enumerate(asks):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(14 if i else 0)
        run(p, f"{i + 1}.  ", size=15, bold=True, color=AMBER)
        run(p, t, size=15, color=BODY)

    footer(s, "06")
    return s


def slide_close(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=WHITE)
    header(s, "05  ·  CONFIRM", "Tracker close-out, and the decisions we need today")

    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 1.18, 6.05, 3.55,
          fill=GREEN_BG, line="B7E0C5")
    tf = tb(s, 0.70, 1.35, 5.55, 0.36)
    p0(tf, "Please confirm these can close", size=14, bold=True, color=GREEN)
    closes = [
        "Cleaning philosophy — our scheme is issued; remaining comments to be listed or closed.",
        "Operating manual — issued; SW review still outstanding.",
        "Control description — issued; SW review still outstanding.",
        "We will write the replies into the tracker and ask Keith to confirm line by line.",
    ]
    tf = tb(s, 0.70, 1.80, 5.55, 2.70)
    for i, t in enumerate(closes):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(8 if i else 0)
        run(p, "●  ", size=14, color=GREEN)
        run(p, t, size=14, color=BODY)

    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.70, 1.18, 6.15, 3.55,
          fill=NAVY, line=NAVY)
    tf = tb(s, 6.95, 1.35, 5.70, 0.36)
    p0(tf, "Decisions needed before we leave", size=14, bold=True, color="F6C089")
    needs = [
        "Accept the 316L gasket hardness spec — or say what else QA needs.",
        "Mark the three-cavity draft: accepted / change / leave for site.",
        "State whether PIC-028 must be rated for 270 °C.",
        "State whether the current PSV orifice still holds at DN50.",
    ]
    tf = tb(s, 6.95, 1.80, 5.70, 2.70)
    for i, t in enumerate(needs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(8 if i else 0)
        run(p, f"{i + 1}.  ", size=14, bold=True, color="F6C089")
        run(p, t, size=14, color=WHITE)

    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 4.90, 12.4, 2.00,
          fill=ICE, line=LINE)
    tf = tb(s, 0.70, 5.05, 11.90, 0.30)
    p0(tf, "Not for this meeting", size=13, bold=True, color=MUTED)
    tf = tb(s, 0.70, 5.38, 11.90, 1.40)
    p0(tf, "PFD and mass balance stay unchanged.  Reactor size stays a September data decision (DN65 working case).  HAZOP and the Sandwich week are a separate scheduling discussion — we are targeting mid-October, subject to Keith and Claire.",
       size=15, color=BODY)
    addp(tf, "Please do not leave open technical points “to finish after we get home”.  If a point cannot close today, write the remaining action and the owner in the tracker before we stop.",
         size=15, color=BODY, before=8)

    footer(s, "07")
    return s


def main():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide_title(prs)
    slide_agenda(prs)
    slide_gasket(prs)
    slide_layout(prs)
    slide_valve(prs)
    slide_psv(prs)
    slide_close(prs)
    prs.save(OUT)
    print("Saved", OUT, "slides", len(prs.slides))


if __name__ == "__main__":
    main()
